"""Exportacion atomica; arrays NPZ sin pickle y metadatos JSON independientes."""
import hashlib
import json
import os
from pathlib import Path
from uuid import uuid4
import numpy as np
import pandas as pd


def jsonable(value):
    if isinstance(value, Path): return str(value)
    if isinstance(value, np.ndarray): return jsonable(value.tolist())
    if isinstance(value, np.generic): return jsonable(value.item())
    if isinstance(value, float) and not np.isfinite(value): return None
    if isinstance(value, dict): return {str(k):jsonable(v) for k,v in value.items()}
    if isinstance(value, (tuple,list)): return [jsonable(v) for v in value]
    return value


def write_json(path, value):
    path=Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    tmp=path.with_name(path.name+'.'+uuid4().hex+'.tmp')
    try:
        tmp.write_text(json.dumps(jsonable(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
        os.replace(tmp,path)
    finally:
        tmp.unlink(missing_ok=True)


def sha256_file(path):
    h=hashlib.sha256()
    with open(path,'rb') as stream:
        for chunk in iter(lambda:stream.read(8*1024*1024),b''): h.update(chunk)
    return h.hexdigest()


def save_details(base, result, metadata):
    """JSON describe rutas de arrays y dimensiones; NPZ se abre con allow_pickle=False."""
    arrays={}
    def pack(value, key):
        if isinstance(value,np.ndarray):
            if value.dtype.hasobject: return pack(value.tolist(),key)
            arrays[key]=value
            return dict(npz_key=key, shape=list(value.shape), dtype=str(value.dtype))
        if isinstance(value,dict): return {str(k):pack(v,key+'__'+str(k)) for k,v in value.items()}
        if isinstance(value,(list,tuple)):
            return [pack(v,key+'__'+str(i)) for i,v in enumerate(value)]
        return jsonable(value)
    structure=pack(result,'result')
    base=Path(base); base.parent.mkdir(parents=True,exist_ok=True)
    target=base.with_suffix('.npz'); tmp=target.with_name(target.name+'.tmp')
    try:
        with open(tmp,'wb') as stream: np.savez_compressed(stream,**arrays)
        os.replace(tmp,target)
    finally: tmp.unlink(missing_ok=True)
    write_json(base.with_suffix('.json'),dict(metadata=metadata,results=structure))


def quality_rows(result, common):
    rows=[]
    for key,value in result.items():
        if isinstance(value,np.ndarray) and np.issubdtype(value.dtype,np.number):
            finite=np.isfinite(value)
            row={**common,'array':key,'shape':str(value.shape),'n_values':value.size,
                 'n_finite':int(finite.sum()),'n_nan':int(np.isnan(value).sum()),
                 'n_inf':int(np.isinf(value).sum()),'note':''}
            if key.startswith('te_'): row['note']='La diagonal de auto-TE es NaN por definicion.'
            rows.append(row)
    return rows


def export_workbook(path, rows, trace, intervals, issues, quality, regions, csv=False):
    path=Path(path); tmp=path.with_name(path.stem+'.'+uuid4().hex+'.tmp.xlsx')
    table=pd.DataFrame(rows)
    if not table.empty:
        ids=['subject','band','condition']
        if table.duplicated(ids).any(): raise ValueError('Claves duplicadas en tabla agregada.')
        table=table[ids+sorted(c for c in table if c not in ids)]
    sheets={'subject_level':table,'trazabilidad':pd.DataFrame(trace),
            'intervalos':pd.DataFrame(intervals),'incidencias':pd.DataFrame(issues),
            'calidad':pd.DataFrame(quality),'regiones':pd.DataFrame(regions)}
    try:
        with pd.ExcelWriter(tmp,engine='openpyxl') as writer:
            for name,df in sheets.items():
                df=df.copy()
                for column in df.columns:
                    df[column]=df[column].map(lambda v:json.dumps(jsonable(v),ensure_ascii=False)
                        if isinstance(v,(dict,list,tuple)) else v)
                df.to_excel(writer,sheet_name=name,index=False)
                ws=writer.sheets[name]; ws.freeze_panes='A2'
                if len(df.columns):
                    ws.auto_filter.ref=ws.dimensions
                    from openpyxl.styles import Font, PatternFill
                    from openpyxl.utils import get_column_letter
                    for cell in ws[1]:
                        cell.font=Font(bold=True,color='FFFFFF')
                        cell.fill=PatternFill('solid',fgColor='215968')
                    for i,column in enumerate(df.columns,1):
                        ws.column_dimensions[get_column_letter(i)].width=min(45,max(14,len(str(column))+2))
        os.replace(tmp,path)
    finally: tmp.unlink(missing_ok=True)
    if csv:
        target=path.with_suffix('.csv'); tmp_csv=target.with_suffix('.csv.tmp')
        try:
            table.to_csv(tmp_csv,index=False,sep=',',decimal='.',encoding='utf-8-sig',float_format='%.17g')
            os.replace(tmp_csv,target)
        finally: tmp_csv.unlink(missing_ok=True)


# Ejes de los arrays que se exportan; el orden de nombres esta en results.
ARRAY_AXES = {
    'theta_power_channels': ['window','channel'],
    'aperiodic_exponent_channels': ['window','channel'],
    'aperiodic_offset_channels': ['window','channel'],
    'theta_power_zones': ['window','region'],
    'aperiodic_exponent_zones': ['window','region'],
    'aperiodic_offset_zones': ['window','region'],
    'wpli_trials': ['block','region','region'],
    'wpli_channels': ['block','channel','channel'],
    'wsmi_matrix': ['region','region','window'],
    'wsmi_channels': ['window','channel','channel'],
    'pe_matrix_channels': ['channel','window'],
    'pe_matrix_zones': ['region','window'],
    'lzc_matrix_channels': ['channel','window'],
    'lzc_matrix_zones': ['region','window'],
    'te_full': ['window','source_region','target_region','lag'],
    'te_mean_lag': ['window','source_region','target_region'],
}


def contribution_rows(result, common):
    """Numero de ventanas/bloques finitos que contribuyen a cada feature regional."""
    rows=[]
    for key in ['theta_power_zones','aperiodic_exponent_zones','aperiodic_offset_zones',
                'wpli_trials','wsmi_matrix','pe_matrix_zones','lzc_matrix_zones','te_mean_lag']:
        if key not in result: continue
        value=np.asarray(result[key]); axes=ARRAY_AXES[key]
        if value.ndim!=len(axes): continue
        unit='block' if 'block' in axes else 'window'
        axis=axes.index(unit); array=np.moveaxis(value,axis,0)
        names=result['zone_names']
        for idx in np.ndindex(array.shape[1:]):
            if key=='te_mean_lag' and idx[0]==idx[1]: continue
            values=array[(slice(None),)+idx]
            rows.append(dict(**common,array=key,scope='aggregate_contributors',unit=unit,
                region_or_pair=' -> '.join(names[i] for i in idx),n_values=values.size,
                n_finite=int(np.isfinite(values).sum()),n_nan=int(np.isnan(values).sum()),
                n_inf=int(np.isinf(values).sum())))
    return rows
