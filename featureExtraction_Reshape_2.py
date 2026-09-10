"""Transforma la tabla agregada del pipeline de formato largo a ancho.

El script esta pensado para ejecutarse directamente desde VS Code. Las rutas de
entrada y salida se editan en la seccion CONFIGURACION; no usa argumentos de
linea de comandos.
"""

from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# =============================================================================
# CONFIGURACION: editar estas rutas antes de ejecutar
# =============================================================================

INPUT_FILE = Path(
    r"D:\Doctorado\protocol2023\resultados_pipeline"
    r"\20260910T145431Z_7b82c740\EEG_features_subject_level.xlsx"
)
OUTPUT_FILE = Path(
    r"D:\Doctorado\protocol2023\resultados_pipeline"
    r"\20260910T145431Z_7b82c740\EEG_features_wide.xlsx"
)

SOURCE_SHEET = "subject_level"
EXPECTED_BANDS = ("delta", "theta", "alpha", "beta")

# Estas variables se calculan en banda ancha o sobre el mismo ajuste 1-30 Hz y
# aparecen repetidas en cada fila de banda. En la salida quedan una sola vez.
SHARED_PREFIXES = ("spec_exp_", "spec_off_", "te_", "lzc_")

# Tolerancia para diferencias numericas insignificantes entre copias de una
# variable compartida. Una mezcla de valores finitos y NaN siempre es un error.
SHARED_RTOL = 1e-10
SHARED_ATOL = 1e-12


ID_COLS = ["subject", "condition"]
BAND_COL = "band"
RESERVED_OUTPUT_SHEETS = {
    "subject_level",
    "mapeo_columnas",
    "validacion_reshape",
    "bandas_faltantes",
    "controles_compartidas",
}


def _validate_paths(input_file: Path, output_file: Path) -> None:
    if not input_file.is_file():
        raise FileNotFoundError(
            f"No existe el archivo de entrada:\n{input_file}")
    if input_file.resolve() == output_file.resolve():
        raise ValueError(
            "INPUT_FILE y OUTPUT_FILE deben ser archivos distintos.")
    output_file.parent.mkdir(parents=True, exist_ok=True)


def _read_source(input_file: Path) -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    with pd.ExcelFile(input_file) as workbook:
        if SOURCE_SHEET not in workbook.sheet_names:
            raise ValueError(
                f"No existe la hoja {SOURCE_SHEET!r}. "
                f"Hojas disponibles: {workbook.sheet_names}"
            )

        data = pd.read_excel(
            workbook,
            sheet_name=SOURCE_SHEET,
            dtype={"subject": "string", "band": "string"},
        )
        supporting = {
            sheet: pd.read_excel(workbook, sheet_name=sheet)
            for sheet in workbook.sheet_names
            if sheet != SOURCE_SHEET and sheet not in RESERVED_OUTPUT_SHEETS
        }
    return data, supporting


def _validate_and_prepare(data: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    required = ID_COLS + [BAND_COL]
    missing_columns = [
        column for column in required if column not in data.columns]
    if missing_columns:
        raise ValueError(f"Faltan columnas identificadoras: {missing_columns}")
    if data.empty:
        raise ValueError("La hoja subject_level esta vacia.")
    if data[required].isna().any().any():
        bad_rows = data.index[data[required].isna().any(axis=1)].tolist()
        raise ValueError(
            f"Hay identificadores o bandas vacios en las filas: {bad_rows[:20]}")

    prepared = data.copy()
    prepared[BAND_COL] = prepared[BAND_COL].str.strip().str.lower()

    unexpected = sorted(set(prepared[BAND_COL]) - set(EXPECTED_BANDS))
    if unexpected:
        raise ValueError(
            f"Hay bandas no configuradas: {unexpected}. "
            "Revisar EXPECTED_BANDS antes de continuar."
        )

    duplicate_mask = prepared.duplicated(ID_COLS + [BAND_COL], keep=False)
    if duplicate_mask.any():
        duplicate_keys = (
            prepared.loc[duplicate_mask, ID_COLS + [BAND_COL]]
            .sort_values(ID_COLS + [BAND_COL])
            .to_dict("records")
        )
        raise ValueError(
            "Hay filas duplicadas para subject-condition-band. "
            f"Ejemplos: {duplicate_keys[:20]}"
        )

    feature_cols = [
        column for column in prepared.columns if column not in required]
    if not feature_cols:
        raise ValueError("No se encontraron columnas de features.")

    # Evita que texto accidental entre a la normalizacion o al PCA.
    converted = prepared[feature_cols].apply(pd.to_numeric, errors="coerce")
    invalid = prepared[feature_cols].notna() & converted.isna()
    if invalid.any().any():
        examples = []
        for row, column in zip(*np.where(invalid.to_numpy())):
            examples.append(
                {
                    "fila": int(prepared.index[row]),
                    "columna": feature_cols[column],
                    "valor": prepared.iloc[row][feature_cols[column]],
                }
            )
            if len(examples) == 20:
                break
        raise ValueError(
            f"Hay valores no numericos entre las features: {examples}")
    prepared[feature_cols] = converted

    return prepared, feature_cols


def _check_shared_columns(
    data: pd.DataFrame, shared_cols: list[str]
) -> pd.DataFrame:
    checks: list[dict] = []
    conflicts: list[dict] = []

    for column in shared_cols:
        groups_checked = 0
        largest_difference = 0.0
        for key, group in data.groupby(ID_COLS, sort=False, dropna=False):
            values = group[column]
            n_missing = int(values.isna().sum())
            n_present = int(values.notna().sum())
            groups_checked += 1

            if n_missing and n_present:
                conflicts.append(
                    {
                        "subject": key[0],
                        "condition": key[1],
                        "feature": column,
                        "problema": "mezcla de valores numericos y NaN entre bandas",
                    }
                )
                continue

            finite = values.dropna().to_numpy(dtype=float)
            if finite.size > 1:
                difference = float(np.max(np.abs(finite - finite[0])))
                largest_difference = max(largest_difference, difference)
                if not np.allclose(
                    finite, finite[0], rtol=SHARED_RTOL, atol=SHARED_ATOL
                ):
                    conflicts.append(
                        {
                            "subject": key[0],
                            "condition": key[1],
                            "feature": column,
                            "problema": f"valores diferentes entre bandas: {finite.tolist()}",
                        }
                    )

        checks.append(
            {
                "feature": column,
                "grupos_verificados": groups_checked,
                "max_diferencia_absoluta": largest_difference,
                "rtol": SHARED_RTOL,
                "atol": SHARED_ATOL,
                "estado": "ok" if not any(c["feature"] == column for c in conflicts) else "error",
            }
        )

    if conflicts:
        raise ValueError(
            "Las variables compartidas no son constantes dentro de "
            "subject-condition. No se genero la salida. Ejemplos: "
            f"{conflicts[:20]}"
        )

    return pd.DataFrame(checks)


def _expected_index(data: pd.DataFrame) -> tuple[pd.MultiIndex, pd.DataFrame]:
    pairs = data[ID_COLS].drop_duplicates().reset_index(drop=True)
    bands = pd.DataFrame({BAND_COL: EXPECTED_BANDS})
    expected = pairs.merge(bands, how="cross")
    actual = data[ID_COLS + [BAND_COL]].drop_duplicates()

    missing_bands = expected.merge(
        actual,
        on=ID_COLS + [BAND_COL],
        how="left",
        indicator=True,
    )
    missing_bands = missing_bands.loc[
        missing_bands["_merge"] == "left_only", ID_COLS + [BAND_COL]
    ].reset_index(drop=True)

    return pd.MultiIndex.from_frame(expected), missing_bands


def reshape_features(
    data: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    prepared, feature_cols = _validate_and_prepare(data)
    shared_cols = [
        column for column in feature_cols if column.startswith(SHARED_PREFIXES)
    ]
    banded_cols = [
        column for column in feature_cols if column not in shared_cols]

    shared_checks = _check_shared_columns(prepared, shared_cols)
    expected_index, missing_bands = _expected_index(prepared)
    pair_index = expected_index.droplevel(BAND_COL).drop_duplicates()

    if shared_cols:
        shared = (
            prepared[ID_COLS + shared_cols]
            .drop_duplicates()
            .set_index(ID_COLS)
            .reindex(pair_index)
        )
    else:
        shared = pd.DataFrame(index=pair_index)

    if banded_cols:
        banded_long = prepared.set_index(ID_COLS + [BAND_COL])[banded_cols]
        # Reindexar crea explicitamente las bandas faltantes con NaN.
        banded_long = banded_long.reindex(expected_index)
        banded_wide = banded_long.unstack(BAND_COL)
        ordered_columns = pd.MultiIndex.from_product(
            [banded_cols, EXPECTED_BANDS], names=["feature", BAND_COL]
        )
        banded_wide = banded_wide.reindex(columns=ordered_columns)
        banded_wide.columns = [
            f"{feature}__{band}" for feature, band in banded_wide.columns
        ]
    else:
        banded_wide = pd.DataFrame(index=pair_index)

    wide = pd.concat([shared, banded_wide], axis=1).reset_index()
    wide = wide.sort_values(ID_COLS, kind="stable").reset_index(drop=True)

    mapping_rows = [
        {
            "columna_entrada": column,
            "banda": "compartida",
            "columna_salida": column,
            "clasificacion": "sin sufijo de banda",
        }
        for column in shared_cols
    ]
    mapping_rows.extend(
        {
            "columna_entrada": column,
            "banda": band,
            "columna_salida": f"{column}__{band}",
            "clasificacion": "dependiente de banda",
        }
        for column in banded_cols
        for band in EXPECTED_BANDS
    )
    mapping = pd.DataFrame(mapping_rows)

    validation = pd.DataFrame(
        [
            {"control": "filas_entrada", "valor": len(prepared)},
            {"control": "filas_salida", "valor": len(wide)},
            {"control": "pares_subject_condition", "valor": len(pair_index)},
            {"control": "bandas_esperadas",
                "valor": ", ".join(EXPECTED_BANDS)},
            {"control": "filas_de_banda_faltantes",
                "valor": len(missing_bands)},
            {"control": "features_compartidas", "valor": len(shared_cols)},
            {"control": "features_dependientes_de_banda",
                "valor": len(banded_cols)},
            {"control": "claves_duplicadas", "valor": 0},
        ]
    )

    return wide, mapping, validation, missing_bands, shared_checks


def _format_sheet(worksheet, dataframe: pd.DataFrame) -> None:
    worksheet.freeze_panes = "A2"
    if len(dataframe.columns):
        worksheet.auto_filter.ref = worksheet.dimensions
        fill = PatternFill("solid", fgColor="215968")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill

        for position, column in enumerate(dataframe.columns, start=1):
            values = dataframe[column].astype(str).head(200)
            content_width = max(
                [len(str(column)), *(len(value) for value in values)])
            worksheet.column_dimensions[get_column_letter(position)].width = min(
                50, max(12, content_width + 2)
            )


def _write_output(
    output_file: Path,
    wide: pd.DataFrame,
    mapping: pd.DataFrame,
    validation: pd.DataFrame,
    missing_bands: pd.DataFrame,
    shared_checks: pd.DataFrame,
    supporting_sheets: dict[str, pd.DataFrame],
) -> None:
    temporary = output_file.with_name(f"{output_file.stem}.tmp.xlsx")
    sheets = {
        "subject_level": wide,
        "mapeo_columnas": mapping,
        "validacion_reshape": validation,
        "bandas_faltantes": missing_bands,
        "controles_compartidas": shared_checks,
        **supporting_sheets,
    }

    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for sheet_name, dataframe in sheets.items():
                dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
                _format_sheet(writer.sheets[sheet_name], dataframe)
        os.replace(temporary, output_file)
    finally:
        if temporary.exists():
            temporary.unlink()


def main() -> None:
    _validate_paths(INPUT_FILE, OUTPUT_FILE)
    data, supporting_sheets = _read_source(INPUT_FILE)
    wide, mapping, validation, missing_bands, shared_checks = reshape_features(
        data)
    _write_output(
        OUTPUT_FILE,
        wide,
        mapping,
        validation,
        missing_bands,
        shared_checks,
        supporting_sheets,
    )

    print(f"Archivo cargado: {INPUT_FILE}")
    print(f"Filas de entrada: {len(data)}")
    print(f"Filas y columnas de salida: {wide.shape}")
    print(f"Bandas faltantes registradas: {len(missing_bands)}")
    print(f"Archivo exportado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
